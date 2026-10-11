"""Builds docs/Complete_Guide.docx: a from-zero explanation of the whole project for a complete beginner.
Usage: python scripts/build_guide.py"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
D, S = ROOT / "docs" / "diagrams", ROOT / "docs" / "screenshots"
NAVY = RGBColor(0x1E, 0x3A, 0x5F)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.95)
sec.top_margin = sec.bottom_margin = Inches(0.85)
doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(11)
for n, sz in (("Heading 1", 20), ("Heading 2", 14.5), ("Heading 3", 12)):
    h = doc.styles[n]
    h.font.name, h.font.size, h.font.color.rgb = "Calibri", Pt(sz), NAVY


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def H1(t):
    doc.add_page_break()
    doc.add_heading(t, 1)


def H2(t):
    doc.add_heading(t, 2)


def P(t, bold=False, size=None, color=None, align=None, after=6, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    return p


def B(items, style="List Bullet"):
    for t in items:
        p = doc.add_paragraph(style=style)
        if t.startswith("**") and "**" in t[2:]:
            head, rest = t[2:].split("**", 1)
            r = p.add_run(head); r.bold = True
            p.add_run(rest)
        else:
            p.add_run(t)


def N(items):
    B(items, "List Number")


def CODE(t):
    tb = doc.add_table(rows=1, cols=1)
    tb.style = "Table Grid"
    c = tb.rows[0].cells[0]
    shade(c, "F3F4F6")
    c.text = ""
    for i, line in enumerate(t.strip("\n").split("\n")):
        para = c.paragraphs[0] if i == 0 else c.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        r = para.add_run(line)
        r.font.name, r.font.size = "Consolas", Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def NOTE(title, text, fill="E8F1FA"):
    tb = doc.add_table(rows=1, cols=1)
    tb.style = "Table Grid"
    c = tb.rows[0].cells[0]
    shade(c, fill)
    c.text = ""
    r = c.paragraphs[0].add_run(title + "  "); r.bold = True; r.font.size = Pt(10.5)
    r2 = c.paragraphs[0].add_run(text); r2.font.size = Pt(10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def FIG(path, cap, w=6.5):
    doc.add_picture(str(path), width=Inches(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(cap, italic=True, size=9.5, align="c", color=RGBColor(0x47, 0x55, 0x69), after=10)


def TABLE(headers, rows, widths=None, font=9.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(font); r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1E3A5F")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(v)).font.size = Pt(font)
    if widths:
        for i, w in enumerate(widths):
            t.columns[i].width = Inches(w)
        for row in t.rows:
            for c, w in zip(row.cells, widths):
                c.width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# =============================================================================== cover
P("PIMPRI CHINCHWAD UNIVERSITY", True, 16, NAVY, "c", 0)
P("School of Computer Applications · MCA", size=11, align="c", after=40)
P("COMPLETE GUIDE", True, 12, RGBColor(0x1F, 0x8A, 0x8A), "c", 8)
P("CURIO", True, 40, NAVY, "c", 0)
P("How the whole project works, explained from zero", italic=True, size=15, align="c", after=36)
P("A luxury e-commerce website with explainable recommendations", size=12, align="c", after=40)
TABLE(["", ""], [["Who this guide is for", "Someone who has never built a website. No prior knowledge is assumed."],
                 ["What you will learn", "How every part works: the page you see, the Python behind it, the database, the recommendations, the payments, the tests and how to run it."],
                 ["How to read it", "Chapters 1-3 give the big picture. Chapters 4-18 explain each part. Chapter 19 answers the hard technical questions an examiner may ask."],
                 ["Project", "Python · Streamlit · MongoDB · scikit-learn · Docker"]], [1.8, 4.6], 10.5)

doc.add_page_break()
doc.add_heading("Contents", 1)
for line in ["1  What is this project, in one page", "2  Words you need first (glossary)", "3  The big picture: stack and architecture", "4  A tour of the folders and files",
             "5  How the page you see is made (Streamlit)", "6  The look and the animations (CSS)", "7  The database (MongoDB) and the data dictionary",
             "8  Follow one click: opening a product", "9  Recommendation 1: 'You may also like' (TF-IDF)", "10  Recommendation 2: 'Frequently bought together'", "11  Accounts and security",
             "12  Bag, coupons, delivery and tax", "13  Orders and stock", "14  Payments (the demo Razorpay)", "15  The admin console and its numbers",
             "16  Where the real data and photos came from", "17  Running it: Docker, start.sh and stop.sh", "18  Testing: how we know it works",
             "19  Hard questions and clear answers", "20  Limits and future work", "21  One-page cheat sheet", "Appendix  UML and data-flow diagrams"]:
    P(line, after=2)

# =============================================================================== 1
H1("1  What is this project, in one page")
P("CURIO is a website that sells luxury goods: watches, handbags, jewellery, perfume, sunglasses, clothing, shoes and fine leather and writing items. "
  "You can browse, search, save favourites, put things in a bag, pay (with a pretend payment gateway) and track your order. A shop owner can log in to an admin console "
  "to see sales charts and manage products and orders.")
P("The special part, and the reason this is a recommendation project, is the suggestions. When you open a product the page shows two kinds:")
B(["**You may also like**: other pieces whose descriptions are similar. A computer finds them with a technique called TF-IDF and cosine similarity (Chapter 9).",
   "**Frequently bought together**: pieces that customers tend to buy in the same order, found by counting past orders (Chapter 10)."])
P("Every suggestion comes with a reason you can read, for example the similarity score or 'In 31% of orders with this piece'. That is what 'explainable' means: the system can show why it suggested something.")
H2("What is real and what is pretend")
TABLE(["Part", "Real or simulated?", "Note"],
      [["Products, brands, models", "Real maisons and real models (Rolex Submariner, Hermès Kelly ...)", "Prices are indicative, not live"],
       ["Descriptions", "30 are real Wikipedia text; the rest are short factual descriptions written for the demo", "Each product stores where its text came from"],
       ["Photographs", "Real professional photos from Unsplash", "They show the piece or a very similar one"],
       ["Payments", "Simulated, modelled on Razorpay test mode", "No money moves; no card is stored"],
       ["Past orders used for suggestions", "Synthetic (generated), flagged in the database", "Real orders you place are added to the same pool"],
       ["Delivery, tax, refunds", "Simulated", "Nothing is shipped"]], [1.8, 2.7, 2.0])
FIG(S / "11_product_page.png", "Figure 1.1 - The product page. The big image is on the left; details stay in view on the right as you scroll.", 5.8)

# =============================================================================== 2
H1("2  Words you need first (glossary)")
P("Skim this chapter now and come back whenever a word is unfamiliar. Each word is explained in everyday language first.")
TABLE(["Word", "Plain-language meaning"],
      [["Browser", "The program you use to visit websites (Chrome, Safari). It draws pages and sends your clicks."],
       ["Server", "A computer that runs the website's program and answers the browser. Here, your own laptop plays the server."],
       ["localhost", "A special address meaning 'this very computer'. http://localhost:8501 is the website running on your machine, on port 8501."],
       ["Port", "A numbered door on a computer. Different programs listen on different doors (the website on 8501, the database on 27017)."],
       ["Frontend / backend", "Frontend is what you see. Backend is the hidden part that stores data and applies rules. Here both are written in Python."],
       ["Framework", "A ready-made toolkit that handles common work. Streamlit is a framework that builds web pages from Python."],
       ["Database", "A program that stores data safely on disk and finds it quickly. We use MongoDB."],
       ["Document database", "MongoDB stores each record as a flexible 'document' (like a JSON object) instead of a strict table row."],
       ["Collection", "A group of similar documents in MongoDB, like a table (users, products, orders ...)."],
       ["Index", "A lookup structure that makes searches fast and can enforce rules such as 'no two users share an email'."],
       ["CSS", "The language that styles web pages: fonts, colours, layout, animation."],
       ["Hash", "A one-way scramble of data. You can check a password against its hash but cannot turn the hash back into the password."],
       ["Salt", "Random extra data mixed into a password before hashing so two equal passwords get different hashes."],
       ["HMAC", "A fingerprint made from a message plus a secret key. Only someone with the key can produce a matching one."],
       ["TF-IDF", "A way to turn text into numbers by counting words and giving rare words more weight than common ones."],
       ["Cosine similarity", "A score from 0 to 1 for how alike two lists of numbers are. 1 means the same direction, 0 means nothing in common."],
       ["Confidence / lift", "Two numbers describing how strongly two products are bought together (Chapter 10)."],
       ["Docker", "A tool that runs a program inside a sealed box (a container) so it behaves the same on every computer."],
       ["Session", "The memory the app keeps about your visit (who you are, what page you are on)."],
       ["Validation", "Checking that what a user typed makes sense before using it (a real email, a positive quantity ...)."],
       ["Test", "A small program that checks another program behaves correctly, automatically and repeatably."],
       ["mongomock", "A fake in-memory MongoDB used when no real database is available, and in fast tests."]], [1.5, 5.0])

# =============================================================================== 3
H1("3  The big picture: stack and architecture")
P("A 'stack' is the set of technologies layered on top of each other. Read Figure 3.1 from top to bottom: your click travels down the layers and the answer travels back up.")
FIG(D / "14_tech_stack.png", "Figure 3.1 - The technology stack. Each layer only talks to the layer next to it.", 6.2)
H2("The same idea as an architecture diagram")
FIG(D / "02_architecture.png", "Figure 3.2 - System architecture: presentation, application services, data and recommendation engines.", 6.0)
H2("Why these choices")
TABLE(["Choice", "Why"],
      [["Python everywhere", "One language for the page, the rules, the maths and the tests. Easy to read and the standard language for data work."],
       ["Streamlit", "Builds a real web page from Python, so we do not need a separate JavaScript front-end. Good fit for a student project that must also be correct and testable."],
       ["MongoDB", "The synopsis asked for it. A product has very different attributes (a watch vs a perfume), which suits flexible documents. Orders also naturally contain their lines."],
       ["scikit-learn", "Trusted library for TF-IDF and cosine similarity, so we do not write the maths by hand."],
       ["Docker for MongoDB only", "Installing a database by hand is error-prone. One command gives everyone the same MongoDB."],
       ["CSS only for animation", "No extra JavaScript to maintain; the browser does the work smoothly."]], [1.8, 4.7])
H2("The three-layer rule used in the code")
P("The code follows a simple rule that keeps it understandable: pages never touch the database directly.")
B(["**ui/** (pages): decides what is drawn and what a button does. It only calls services.",
   "**shop/** (services): holds the rules. 'Can this quantity be added?', 'Is this coupon valid?', 'Reserve the stock'. It is the only code that talks to MongoDB.",
   "**MongoDB**: stores the data."])
P("Because rules live in shop/, they can be tested without opening a browser, and the same rules protect every page.")

# =============================================================================== 4
H1("4  A tour of the folders and files")
CODE("""project/
  app.py                 starting point: sets the page, loads the style, shows header + current page + footer
  ui/                    everything you SEE (pages)
    header.py            fixed top bar + footer
    landing.py           the animated home page
    shop_pages.py        collection page and product page
    components.py        the product 'tile' and the grid
    checkout.py          bag, checkout, payment dialog, order tracking, confirmation
    account.py           log in / register, orders, wishlist, addresses, profile
    admin.py             admin console (dashboard, orders, products, coupons)
    help_page.py         delivery / payments / test credentials
    state.py             wiring: creates services once, holds shared helpers and session defaults
    icons.py             small SVG icons and star ratings
  shop/                  everything the site DOES (rules)
    db.py                connects to MongoDB (or the in-memory fake)
    config.py            settings (currency, GST, shipping fees ...)
    validators.py        input checks
    users.py             accounts, password hashing, roles
    catalog.py           search, filters, sorting, product add/edit/delete
    cart.py coupons.py pricing.py orders.py payments.py     the buying flow
    recommender.py       'You may also like' (TF-IDF)
    bundles.py           'Frequently bought together' (basket analysis)
    reviews.py wishlist.py addresses.py stats.py invoice.py
    seed_data.py synthetic.py catalog_data.json brands.py   the starting data
    theme.py             all the CSS (the look and the animation)
  tests/                 automated checks (158 tests)
  scripts/               one-off tools: build the catalogue, fetch photos, make diagrams, evaluate
  assets/                product photos (products/) and landing photos (hero/) + CREDITS.md
  docs/                  diagrams, screenshots, this guide and the project report
  start.sh stop.sh       one-command start and stop
  docker-compose.yml     describes the MongoDB container""")

# =============================================================================== 5
H1("5  How the page you see is made (Streamlit)")
P("A normal website is made of HTML files. In this project you do not write HTML pages by hand for the shop. You write Python like st.button('Add') and Streamlit sends the browser the right page.")
H2("The rerun idea")
P("The one idea you must understand about Streamlit: every time you click something, Streamlit runs the whole Python script again from the top and redraws the page. "
  "That sounds wasteful but it keeps the code simple, because the page is always a direct picture of the current data.")
FIG(D / "15_streamlit_rerun.png", "Figure 5.1 - The Streamlit loop.", 6.2)
H2("But then how does it remember anything?")
P("It uses st.session_state, a dictionary kept for each browser tab. The project stores things like the logged-in user, the page you are on and the product you opened. A small example:")
CODE("""# ui/state.py
def go(view, selected=None):        # called when you press a button
    ss.view = view                  # remember which page to show
    if selected is not None:
        ss.selected = selected      # remember which product

# app.py
PAGES = {"Home": page_home, "Catalog": page_catalog, "Product": page_product, ...}
PAGES.get(ss.view, page_home)()     # draw the page named in session_state""")
P("So 'navigation' is just changing ss.view and letting the script rerun. There are no separate web addresses for each page.")
H2("Callbacks")
P("A button can run a function before the rerun (a callback). For example the + on a tile calls quick_add(product_id), which asks the cart service to add the item. "
  "Because the callback runs first, by the time the page redraws the bag count is already updated.")
H2("The site map")
FIG(D / "16_site_map.png", "Figure 5.2 - Every page and how you reach it.", 6.0)
H2("Services are created once")
P("Opening a database connection is slow, so ui/state.py builds all the services one time with @st.cache_resource and reuses them for every visitor. This is the 'wiring' file of the project.")

# =============================================================================== 6
H1("6  The look and the animations (CSS)")
P("The design follows the feel of a fashion-house website: very little decoration, small uppercase type, big photographs, black and white. Everything visual is in shop/theme.py as one long CSS string that app.py injects into the page.")
H2("Design rules")
B(["**Type:** Inter (a clean sans-serif) for text and Bodoni Moda (an elegant high-contrast serif) for the wordmark and titles.",
   "**Colour:** black, white and a warm grey behind pictures. Gold appears only in charts.",
   "**No clutter:** no borders on product tiles, no shadows, no rounded corners, no emoji.",
   "**Tiles:** every product picture is 3:4 portrait, so the grid is perfectly even."])
H2("The pinned, colour-inverting wordmark")
P("On the home page the giant word CURIO stays fixed on screen while slides move behind it. Its colour flips depending on what is behind it. This uses a CSS property called mix-blend-mode: difference, which combines the colour of the image behind with the text colour:")
CODE(""".zmark { position: fixed; font-size: 23vw; color: #fff; mix-blend-mode: difference; pointer-events: none; }""")
P("position: fixed keeps it in place; pointer-events: none lets your clicks pass through it to the slide below.")
H2("The scroll 'unfold' animation")
P("Each campaign picture starts as a small framed crop and grows to full screen as you scroll it into view. Modern browsers can tie an animation to scroll position with animation-timeline: view(). No JavaScript is needed:")
CODE("""@keyframes unfold {
  from { clip-path: inset(16% 14% 16% 14%); transform: scale(1.22); }   /* small framed crop */
  to   { clip-path: inset(0);               transform: scale(1);    }   /* full picture */
}
.zs .zi { animation: unfold linear both; animation-timeline: view(); animation-range: entry 0% cover 50%; }""")
FIG(D / "21_scroll_animation.png", "Figure 6.1 - How scroll controls the animation.", 6.0)
P("If a browser does not support it, the rule sits inside @supports (animation-timeline: view()), so the image simply appears normally. "
  "A second rule disables all motion for people who set 'reduce motion' in their system (prefers-reduced-motion).")
H2("The responsive grid")
P("The product grid must show 5 columns on a wide monitor and 2 on a phone. Streamlit puts columns in a row; CSS tells that row to wrap and gives each column a width that depends on the screen width using media queries:")
CODE("""@media (max-width: 700px) { GRIDCOL { flex-basis: calc((100% - var(--gap)) / 2); } }   /* 2 columns on phones */""")
H2("The invisible button trick (and a real bug we fixed)")
P("A product tile should open when you click anywhere on its picture. Streamlit buttons cannot wrap an image, so each tile contains a real button stretched over the picture and made invisible (opacity: 0). "
  "Early on, the button did not stretch to fill the picture, so a real mouse click on the picture did nothing, even though the automated tests passed because they pressed the button directly. "
  "The fix was to force the button and its wrappers to 100% width and height. The lesson: always test the way a person uses the site, with real clicks.", after=8)
NOTE("Why this matters", "A test that presses the hidden button proves the code works. Only a real mouse click on the picture proves the page works for a person. Both are now checked.")

# =============================================================================== 7
H1("7  The database (MongoDB) and the data dictionary")
H2("Tables versus documents")
P("In a traditional database every record is a row with fixed columns. In MongoDB each record is a document, like this order (shortened):")
CODE("""{ "_id": "ORD-20261011-1A2B3C",
  "user_id": "U-demo0001",
  "items": [ {"product_id": "P006", "name": "Submariner", "unit_price": 1130000, "quantity": 1} ],
  "pricing": {"subtotal": 1130000, "discount": 0, "shipping": 0, "gst_included": 172372.88, "total": 1130000},
  "status": "Placed",
  "payment": {"method": "razorpay", "status": "paid"},
  "timeline": [ {"status": "Pending Payment", ...}, {"status": "Placed", ...} ] }""")
P("Notice the order lines, the price breakdown and the timeline live inside the order. One read gives the whole order; this is why a document database suits it.")
H2("The entity relationship diagram (ERD)")
FIG(D / "01_erd.png", "Figure 7.1 - Entities and how they connect (PK = primary key, FK = foreign key).", 6.5)
H2("The nine collections")
TABLE(["Collection", "Holds", "Important rule or index"],
      [["users", "accounts, role, password hash", "email is unique"], ["addresses", "saved delivery addresses", "exactly one default per user"], ["products", "125 pieces", "indexed by category"],
       ["carts", "one cart per user", "user_id unique"], ["wishlists", "saved product ids per user", "-"], ["reviews", "ratings and text", "one review per user per product"],
       ["coupons", "discount rules", "code is the id"], ["orders", "orders with snapshots and timeline", "indexed by user, status and product"], ["payments", "gateway order and payment records", "indexed by order id"]],
      [1.3, 2.7, 2.5])
H2("Data dictionary (the main fields)")
TABLE(["Collection.field", "Type", "Meaning and rule"],
      [["users.email", "string", "Login name. Lower-cased, valid format, must be unique."], ["users.password_hash", "string", "scrypt hash with a random salt. The real password is never stored."],
       ["users.role", "string", "'customer' or 'admin'."], ["products.price", "number", "Indicative retail price in rupees, GST included, above 0."],
       ["products.stock", "integer", "Units available, 0 or more; never negative."], ["products.tags", "list of strings", "Keywords used by the recommender (counted twice)."],
       ["products.source", "string", "Where the description came from (Wikipedia or written for the demo)."], ["products.image", "string", "Relative path to the photo file."],
       ["carts.items[].quantity", "integer", "1 to 10 and not above stock."], ["orders.status", "string", "Pending Payment, Placed, Packed, Shipped, Delivered or Cancelled."],
       ["orders.items[]", "list", "Snapshot of name, brand, price and quantity at the time of purchase."], ["orders.address", "object", "Snapshot of the delivery address."],
       ["orders.synthetic", "boolean", "True for generated history; hidden from revenue by default."], ["payments.amount", "integer", "Amount in paise (1 rupee = 100 paise)."],
       ["payments.signature", "string", "HMAC-SHA256 the server checks."], ["reviews.verified", "boolean", "True if the reviewer has an order containing the product."],
       ["coupons.kind / value", "string / number", "'percent' or 'flat' and its size."]], [1.8, 1.0, 3.7])
P("The complete dictionary with every field is in docs/DATABASE_DESIGN.md.", italic=True)
H2("Why snapshots?")
P("If an order only stored product ids, then raising a price tomorrow would change yesterday's invoice. So when an order is created we copy the name, price and address into it. History never changes.")
H2("Indexes")
P("Without an index MongoDB reads every document to find one. With an index on email it jumps straight to it. A unique index also makes the database itself refuse a second account with the same email, even if two people sign up at the exact same moment.")
H2("Two databases, one code")
P("shop/db.py tries a real MongoDB. If none is running it quietly uses mongomock, an in-memory fake with the same commands. The footer tells you which one is active. Tests use the fake for speed and are also run against the real one.")

# =============================================================================== 8
H1("8  Follow one click: opening a product")
P("Let us follow exactly what happens when you click the Rolex Submariner tile. This is the best way to see how all the parts fit together.")
N(["You click the picture. The invisible button's callback runs go('Product', 'P006'), which stores view='Product' and selected='P006' in session_state.",
   "Streamlit reruns app.py. It draws the header, then calls page_product() because session_state says 'Product'.",
   "page_product() asks the ProductRepository for product P006. The repository asks MongoDB (products.find_one) and returns the document.",
   "The page draws the large image (a data-URI made by shop/media.py), the name, price, description and the maison facts from shop/brands.py.",
   "For 'Frequently bought together' it calls BundleService.together('P006'), which reads past orders and returns the top partners with confidence and lift.",
   "For 'You may also like' it calls RecommenderService.recommend('P006'), which compares the product's text vector with all others and returns the best matches with reasons.",
   "Reviews come from ReviewService. The page is sent to the browser and the CSS animation fades it in."])
FIG(D / "08_sequence_recommend.png", "Figure 8.1 - The sequence of calls when a product is opened.", 5.6)
P("Now you press ADD. The callback calls CartService.add(user, product, qty), which validates the quantity, checks stock, then saves the cart document. The bag count in the header updates because the script reran.")

# =============================================================================== 9
H1("9  Recommendation 1: 'You may also like' (TF-IDF)")
P("Goal: given one product, find others that are 'similar'. A computer cannot read, so we must turn words into numbers. We do it in four steps.")
H2("Step 1: build one text per product")
CODE("""text = name + category + brand + tags + tags + description        # tags are repeated so they count double""")
H2("Step 2: TF-IDF turns each text into a list of numbers")
B(["**TF (term frequency):** how often a word appears in this product's text. 'watch' appearing twice means it matters here.",
   "**IDF (inverse document frequency):** a word that appears in almost every product ('the', 'luxury') tells us little, so it gets a low weight. A rare word ('chronograph') gets a high weight.",
   "Common English words (a, the, of) are removed (stop words). sublinear_tf squashes very high counts so one repeated word cannot dominate."])
H2("Step 3: cosine similarity compares two lists")
P("Picture each product as an arrow in a space with one direction per word. Two arrows pointing the same way are alike. Cosine similarity measures the angle: 1 means identical direction, 0 means nothing in common. Because TF-IDF rows are scaled to length 1, cosine similarity is just a dot product, which is very fast.")
FIG(D / "17_tfidf_example.png", "Figure 9.1 - A tiny worked example with three products.", 6.4)
H2("Step 4: filter and explain")
B(["Remove the product itself, products with the same name, and sold-out products.", "Take the top N (4 or 8).",
   "If nothing is similar, fall back to the same category and say so clearly, without inventing a score. If even that is empty, show a clear empty message.",
   "Show reasons by comparing the two products: same category, same brand, shared tags."])
FIG(D / "11_flow_recommendation.png", "Figure 9.2 - The recommendation pipeline as a flowchart.", 6.5)
H2("Speed: do the heavy work only when needed")
P("Building the TF-IDF matrix is the slow part. RecommenderService keeps the built model and computes a fingerprint (a SHA-1 of every product's id and text). "
  "It rebuilds only when the fingerprint changes, that is, when a product's text changes. Price and stock changes do not trigger a rebuild; they are applied live to the results.")
H2("How good is it? (the numbers)")
P("We check it with an offline test: for each in-stock product, are the top 5 suggestions in the same subcategory? On this catalogue TF-IDF scores Precision@5 = 0.59 and MRR = 0.84, against 0.40 and 0.56 for a 'same category' baseline and 0.04 and 0.09 for random picks. "
  "This is a proxy, not a measure of sales. See docs/evaluation_results.md for the full table and what the numbers mean.")

# =============================================================================== 10
H1("10  Recommendation 2: 'Frequently bought together'")
P("TF-IDF looks at words. This second method ignores words and only looks at behaviour: which products appear in the same order. It is the same idea as a supermarket noticing that people who buy bread often buy butter.")
H2("Two numbers")
B(["**Confidence(A to B)** = orders containing both A and B divided by orders containing A. 'Of the people who bought the watch, what share also bought the wallet?'",
   "**Lift** = confidence divided by the overall share of orders containing B. It asks: is that better than chance? Lift above 1 means a real link. A very popular item has high confidence with everything, but its lift stays near 1."])
FIG(D / "18_basket_maths.png", "Figure 10.1 - A worked example with 100 orders.", 6.4)
H2("Rules in the code")
B(["Only completed orders count (Placed, Packed, Shipped, Delivered).", "A pair must appear in at least 2 orders (minimum support) so a one-off basket never shows.",
   "Only pairs with lift above 1 are kept; the rest are ranked by confidence x lift.", "The analysis is cached and rebuilt only when the orders change (it checks the number and latest update of orders)."])
H2("Where the starting orders come from")
P("A brand-new shop has no orders, so there is nothing to count. To make the feature demonstrable we generate about 850 past orders from realistic shopping 'missions', for example:")
TABLE(["Mission", "Usually contains"],
      [["Watch collector", "a sports watch, maybe a chronograph, a wallet or belt, a pen"], ["The gift", "a perfume, a necklace, a scarf"], ["Evening", "heels, a handbag, earrings, a perfume"],
       ["Tailoring", "a suit, oxford shoes, a belt, a dress watch"], ["Travel", "luggage, sunglasses, a wallet"]], [1.6, 4.9])
NOTE("About these orders", "These generated orders are flagged synthetic: true in the database, never change stock, are hidden from revenue unless you switch them on, and are disclosed on the page. "
     "Real orders placed in the app are added to the same pool, so the statistics move toward reality over time.", "FFF4E5")

# =============================================================================== 11
H1("11  Accounts and security")
H2("Passwords are never stored")
P("When you register, the password is run through scrypt, a deliberately slow hash, together with a random 16-byte salt. Only the result is stored:")
CODE("scrypt$16384$8$1$<salt in hex>$<hash in hex>")
P("At log-in the same calculation is repeated and the two hashes are compared with a constant-time function (so timing reveals nothing). Slowness is the point: a thief who steals the database must spend a lot of computing power to guess each password.")
B(["Passwords need 8-128 characters with at least one letter and one number.",
   "Wrong email and wrong password give the same message, so an attacker cannot learn which emails exist. Even for an unknown email a dummy hash is computed so the response time is the same.",
   "Roles: 'customer' or 'admin'. The admin console is checked in the page and again inside the service (require_admin), so even a crafted request cannot change an order's status without the role."])
FIG(D / "22_security_layers.png", "Figure 11.1 - The security layers.", 6.0)
H2("Validation")
P("Everything a user types is checked in shop/validators.py before use: names (letters only), emails (format), Indian mobile numbers (10 digits starting 6-9), PIN codes (6 digits, not starting with 0), card numbers (Luhn checksum), UPI ids, quantities (whole number 1-10), search text (length) and reviews (rating 1-5, minimum text length). "
  "Search words are escaped so special characters like ( cannot break the search.")
NOTE("Luhn check", "Card numbers end with a check digit. Doubling every second digit from the right and summing the digits must give a multiple of 10. It catches typing mistakes without contacting a bank.")

# =============================================================================== 12
H1("12  Bag, coupons, delivery and tax")
H2("The bag")
P("One cart document per user. Adding an item checks: quantity is a whole number from 1 to 10, the product exists, it is in stock, and the total in the bag does not exceed stock. If a product is later deleted or stock drops, the bag shows a warning and checkout is blocked until fixed.")
H2("Coupons")
TABLE(["Code", "Rule"], [["WELCOME10", "10% off, minimum order 20,000, capped at 50,000"], ["FLAT5000", "5,000 off orders above 1,00,000"],
                         ["TIMELESS8", "8% off watches and jewellery only, minimum 2,00,000, cap 1,50,000"], ["SCENT15", "15% off fragrance only, minimum 15,000, cap 10,000"]], [1.5, 5.0])
P("A coupon is checked for: existing, active, not expired, under its usage limit, minimum order met, and restricted to certain categories. For category coupons only the eligible lines count towards the discount. "
  "A coupon is marked as used only when payment succeeds (or when a cash-on-delivery order is placed), so abandoned payments do not use it up.")
H2("Delivery and tax")
B(["Insured standard delivery is free above 10,000, otherwise 250. Priority courier is 1,200. Remote PIN codes add 2 days to the estimate.",
   "Prices already include 18% GST. The invoice shows how much of the total is tax: tax = amount x 18 / 118. For 1,18,000 that is 18,000."])
CODE("""total = subtotal - discount + shipping
gst_included = (subtotal - discount) * 0.18 / 1.18""")

# =============================================================================== 13
H1("13  Orders and stock")
H2("The big risk: selling what you do not have")
P("Two customers may buy the last piece at the same moment. The code protects stock with a conditional update that MongoDB performs atomically, meaning it cannot be interrupted halfway:")
CODE("""products.update_one({"_id": id, "stock": {"$gte": qty}},     # only if enough stock remains
                    {"$inc": {"stock": -qty}})                # then subtract in one step
# if modified_count != 1  ->  someone else got it first  ->  raise StockError""")
P("If the second item of an order fails after the first was already reserved, the code adds the first item's quantity back (a rollback), so no stock is lost.")
H2("The life of an order")
FIG(D / "10_order_lifecycle.png", "Figure 13.1 - Order states. Only an admin can move an order forward, one step at a time.", 6.4)
B(["**Pending Payment:** online order created, stock reserved, waiting for payment. It expires after 15 minutes and the stock is released.",
   "**Placed / Packed / Shipped / Delivered:** normal progress. Cash-on-delivery orders become 'paid' when delivered.",
   "**Cancelled:** allowed until the order ships. Stock is restored and, if paid, a refund record is made."])
P("Every change appends an entry to the order's timeline, which is what the tracking page draws.")

# =============================================================================== 14
H1("14  Payments (the demo Razorpay)")
P("Razorpay is a real Indian payment company. Its checkout works in a fixed pattern. This project reproduces that pattern in a safe simulation so you can see and test the whole flow without real money.")
FIG(D / "09_sequence_payment.png", "Figure 14.1 - The online payment sequence.", 5.6)
N(["The order is created as 'Pending Payment' and the server creates a gateway order for the exact amount in paise.",
   "The dialog collects UPI, card or net-banking details. The test gateway accepts only documented test values: UPI success@razorpay pays, failure@razorpay fails; cards 4111 1111 1111 1111 and two others (then a test bank page with Success and Failure buttons).",
   "On success the gateway gives a payment id and a signature.",
   "The browser hands these to the server, which re-computes the signature and checks the amount. Only then is the order marked paid."])
FIG(D / "19_payment_signature.png", "Figure 14.2 - Why the signature check matters.", 6.2)
H2("What is never stored")
P("The card number and CVV are validated and then dropped. Only the network (Visa or Mastercard) and the last four digits are kept.")
H2("Going live")
P("The DemoRazorpay class has the same methods a real integration needs: create_order, sign, verify, authorize. To go live you would implement them with Razorpay's SDK and Checkout.js and add webhooks. The rest of the code would not change.")

# =============================================================================== 15
H1("15  The admin console and its numbers")
P("Log in as admin@curio.test (password Admin@1234). The dashboard reads all orders once and computes everything below.")
FIG(S / "24_admin_dashboard.png", "Figure 15.1 - The admin dashboard.", 6.2)
TABLE(["Number", "How it is calculated"],
      [["Revenue", "Sum of totals of completed orders in the chosen period (cancelled orders are excluded)."], ["Change %", "Compared with the previous period of the same length."],
       ["Average order", "Revenue divided by number of completed orders."], ["Pieces sold", "Sum of quantities over all completed order lines."],
       ["Active buyers", "Distinct customers who ordered in the period. Repeat rate = buyers who also bought in the previous period."],
       ["Inventory value", "Sum of price x stock over all products."], ["Revenue by week", "Orders grouped by the Monday of their week."],
       ["Order status donut", "Count of orders per status."], ["Heat map", "Orders counted by weekday and hour of the day."],
       ["Order value histogram", "Orders grouped into value bands."], ["Low stock", "Products with 3 or fewer pieces."]], [1.7, 4.8])
P("A switch lets you include or exclude the generated history, so you can see either a rich demo or only the orders really placed in the app.")
P("The other tabs: Orders (move an order to Packed, Shipped, Delivered), Products (add, edit, delete and upload a photo with type, size and dimension checks) and Coupons (create, enable, disable).")

# =============================================================================== 16
H1("16  Where the real data and photos came from")
P("The catalogue is not invented. It was built with a repeatable pipeline in scripts/:")
N(["catalog_master.py lists 125 real pieces: maison, model name, collection, an indicative price, stock and tags.",
   "build_catalog.py asks the Wikipedia API for each model's article: the opening paragraph, the page link, and the lead photo if it is freely licensed.",
   "catalog_text.py decides which products use the Wikipedia text (30 where the article is about that very product) and holds short factual descriptions for the rest.",
   "finalize_catalog.py merges everything into shop/catalog_data.json, which is what seeds the database.",
   "unsplash_pick.py searched Unsplash for each brand and model; contact sheets were reviewed by eye and photo_picks.py records the chosen photo for each product.",
   "fetch_product_photos.py downloads the choices, crops them to 3:4 portrait tiles and writes assets/CREDITS.md with every photographer."])
NOTE("Limits you should state clearly", "Official brand packshots are copyrighted and cannot be downloaded, so some photos show a similar piece rather than the exact reference. "
     "Prices are indicative. Brand names belong to their owners; the project is a non-commercial academic demo and is not affiliated with any maison.", "FFF4E5")

# =============================================================================== 17
H1("17  Running it: Docker, start.sh and stop.sh")
H2("Docker in one paragraph")
P("A container is a sealed box that holds a program and everything it needs. docker-compose.yml describes one box: MongoDB 7. 'docker compose up' starts it. A healthcheck repeatedly asks MongoDB 'are you ready?' so we know when it is safe to use. "
  "The database files live in a Docker volume called mongo_data, so your data survives restarts. The port is bound to 127.0.0.1 so only your own computer can reach it.")
H2("./start.sh")
FIG(D / "20_start_script_flow.png", "Figure 17.1 - What start.sh checks, in order.", 5.2)
H2("./stop.sh")
B(["Stops the Streamlit process only if it really is this project's (it checks the command line, so it never kills an unrelated program on the same port).",
   "Removes the MongoDB container but keeps the data volume. Add --purge to delete the data too (it asks you to type 'delete').",
   "Stops the Docker virtual machine only if start.sh started it and nothing else is using it.", "Finally verifies that the port is free and the container is gone."])
H2("Commands you will use")
CODE("""./start.sh               # start everything -> http://localhost:8501
./start.sh --check       # only check prerequisites
./start.sh --mock        # no Docker: use the in-memory database
./stop.sh                # stop the app and MongoDB (data kept)
./stop.sh --purge        # also delete the database volume
pytest -q                # run all 158 tests""")

# =============================================================================== 18
H1("18  Testing: how we know it works")
P("There are three kinds of check.")
TABLE(["Kind", "What it proves", "Examples"],
      [["Service tests", "The rules are right, without any browser", "Coupon limits, stock rollback, tampered payment signature rejected, order cannot be cancelled after shipping, passwords hashed"],
       ["UI journey tests", "Pages work when driven like a user (Streamlit's AppTest)", "Log in, register, add bundle, cash-on-delivery checkout, admin gate, no emoji on the page"],
       ["Real-browser run", "The real site works in real Chrome with real mouse clicks", "scripts/make_screenshots.py: UPI payment, failed payment, card with bank page, COD, cancel, admin; it raises on any surprise"]], [1.4, 2.3, 2.8])
B(["158 tests run on the fast fake database. The service tests (146) and UI tests (12) are also run against a real MongoDB 7 container.",
   "There is a test for each payment failure case: wrong test card, expired card, failed UPI, bank page failure then retry, tampered signature, wrong order id, mismatched amount.",
   "A lesson learned: one early crash happened because a test fixture created hundreds of throw-away databases on the real server. The fixture now reuses one database per run."])

# =============================================================================== 19
H1("19  Hard questions and clear answers")
P("These are questions an examiner may ask, answered in plain language. Short answers first, details after.")
QA = [
 ("Why Streamlit and not React or Django?", "Streamlit lets one language build the page, the rules and the maths. React or Django would need more code and a separate front-end. The trade-off: Streamlit reruns the script on every interaction and gives less control, which we handled with CSS and careful state."),
 ("What is TF-IDF, in one sentence?", "A way to turn text into numbers where words that are frequent in one product but rare across the catalogue get the biggest weights."),
 ("Why cosine similarity and not Euclidean distance?", "Cosine compares direction, so a long description and a short one about the same thing still match. Euclidean distance would be affected by text length."),
 ("Is this machine learning?", "TF-IDF with cosine similarity is a classical, unsupervised information-retrieval method: nothing is trained on labels. Basket analysis is association-rule mining. Both are legitimate data-mining techniques; neither is deep learning."),
 ("What is a content-based recommender and what are its weaknesses?", "It recommends items similar to one you are looking at. Weaknesses: it cannot discover surprising tastes, it depends on good descriptions, and it is not personal to the user."),
 ("Why not collaborative filtering?", "It needs many users with histories. A new store has none. The basket analysis gives a small behaviour-based signal; true collaborative filtering is future work."),
 ("What is cold start?", "The problem of recommending when there is no history. Content-based similarity solves it for products; synthetic starting orders (disclosed) bootstrap the co-purchase feature."),
 ("How do you prevent selling the last item twice?", "A conditional atomic update on stock (only subtract if stock >= quantity) plus rollback of earlier reservations if any line fails. MongoDB performs each update atomically."),
 ("Are your operations transactions?", "Not multi-document transactions. Each stock change is atomic, and failures are compensated by rolling back. For a larger system you would use MongoDB transactions on a replica set."),
 ("How are passwords stored?", "scrypt with a random salt (N=16384, r=8, p=1); only the hash is stored; comparison is constant-time."),
 ("Why a salt?", "So identical passwords produce different hashes and precomputed tables (rainbow tables) are useless."),
 ("What stops a user editing the page to pay less?", "The server never trusts the browser. It recalculates prices from the database and verifies the payment signature and that the paid amount equals the order total."),
 ("What is HMAC?", "A keyed hash. Only someone who knows the secret can produce the right value, so a matching signature proves the data was not altered."),
 ("Is the Razorpay integration real?", "No. It is a faithful simulation of the test-mode flow with the same shape (order, payment id, signature). Switching to the real service means implementing the same methods with Razorpay's SDK."),
 ("Why are card numbers not stored?", "Storing card data brings heavy legal duties (PCI-DSS) and risk. We validate and discard; only network and last four digits remain."),
 ("What happens if payment is abandoned?", "The order stays Pending Payment and holds the stock for 15 minutes; then expire_pending cancels it and returns the stock."),
 ("Why MongoDB?", "The synopsis chose it. Products differ in attributes, and orders contain their lines, which suits documents. A relational database would also work, with more joins."),
 ("What would you change for 1 million products?", "Precompute and store similarity (a nearest-neighbour index), use a search engine for text search, add caching and a replica set, and move the Streamlit UI behind a proper API."),
 ("How does the recommender stay fast?", "Only the selected product's row is multiplied with the sparse matrix (no all-pairs matrix), and the model is rebuilt only when product text changes (content fingerprint)."),
 ("How do you know the recommendations are good?", "An offline evaluation with Precision@K, MRR and coverage against random and category baselines. It is a proxy (same subcategory); real proof needs user clicks or sales."),
 ("What is lift and why use it?", "Lift divides confidence by the item's overall popularity. It stops very popular items from appearing with everything."),
 ("How do you stop one user reading another's order?", "Order lookups include the user id in the query (_id and user_id), so another user's order simply is not found."),
 ("How is the admin protected?", "The role is checked in the UI and again inside the service functions (require_admin), so a crafted call fails with an AuthError."),
 ("What is session_state and what are its limits?", "A per-tab dictionary Streamlit keeps. It is lost when the page is refreshed or the server restarts; a real site would use signed cookies or server-side sessions."),
 ("How is the responsive layout built?", "CSS media queries change the column width of a wrapping flex row: 5, 4, 3 or 2 columns depending on screen width."),
 ("How does the scroll animation work without JavaScript?", "CSS scroll-driven animations: animation-timeline: view() links keyframe progress to how much of the element is visible."),
 ("What is mongomock for?", "An in-memory stand-in for MongoDB: fast tests and a fallback so the app still runs without Docker."),
 ("What are the main limitations?", "Simulated payments and delivery; synthetic starting orders; photos may show similar pieces; indicative prices; item-to-item (not personal) recommendations; single-process Streamlit session state."),
 ("What would you do next?", "Real Razorpay, email and SMS, personalised and hybrid recommendation, proper sessions with cookies, image-based similarity, and a managed cloud database."),
]
for i, (q, a) in enumerate(QA, 1):
    p = doc.add_paragraph()
    r = p.add_run(f"Q{i}. {q}"); r.bold = True
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    doc.add_paragraph(a).paragraph_format.space_after = Pt(6)

# =============================================================================== 20
H1("20  Limits and future work")
B(["Payments, delivery, tax and refunds are simulated. Nothing is charged or shipped.", "The starting 'bought together' orders are generated; they are flagged and disclosed.",
   "Photographs are professional images chosen per brand and model; some show a similar piece. Prices are indicative.",
   "Recommendations are not personal. Sessions are Streamlit session state (no cookies).", "No email or SMS; one store; no multi-document transactions."])
P("Future work:", bold=True)
B(["Real Razorpay (SDK, Checkout.js, webhooks) and a shipping partner.", "Personalised and hybrid recommendation from each customer's browsing and purchase history.",
   "Image-based similarity using computer vision.", "Cloud deployment with a managed database and monitoring.", "A/B tests of recommendation quality with real visitors."])

# =============================================================================== 21
H1("21  One-page cheat sheet")
TABLE(["Question", "Short answer"],
      [["What is it?", "A luxury e-commerce site with explainable recommendations."], ["Language and framework?", "Python 3 with Streamlit."], ["Database?", "MongoDB 7 (PyMongo); mongomock as fallback."],
       ["Method 1?", "TF-IDF + cosine similarity on product text -> You may also like."], ["Method 2?", "Market-basket confidence x lift on past orders -> Frequently bought together."],
       ["How do pages change?", "Buttons change session_state.view; Streamlit reruns the script."], ["How is stock protected?", "Atomic conditional update + rollback."],
       ["How are passwords kept?", "scrypt hash with salt."], ["How are payments verified?", "HMAC-SHA256 signature + amount check on the server."],
       ["How do I run it?", "./start.sh  then open http://localhost:8501"], ["Logins?", "demo@example.com / Demo@1234 ; admin@curio.test / Admin@1234"],
       ["Test payments?", "UPI success@razorpay ; card 4111 1111 1111 1111 (any future expiry, any CVV)"], ["How many tests?", "158, also run on a real MongoDB."],
       ["Where are the diagrams?", "docs/diagrams/ and in this guide and the project report."]], [2.0, 4.5])

# =============================================================================== appendix
H1("Appendix  UML and data-flow diagrams")
P("These are the formal diagrams required for the project. Each is also in the main project report.")
FIG(D / "03_module_hierarchy.png", "Module hierarchy.", 6.6)
FIG(D / "05_use_case.png", "Use case diagram.", 6.5)
FIG(D / "07_class_services.png", "Class diagram (application services).", 6.6)
FIG(D / "12_dfd_level0.png", "Data flow diagram, level 0.", 6.2)
FIG(D / "13_dfd_level1.png", "Data flow diagram, level 1.", 6.5)

out = ROOT / "docs" / "Complete_Guide.docx"
doc.save(out)
print("saved", out)
