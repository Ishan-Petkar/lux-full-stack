"""Builds docs/Viva_Notes.docx: a short, very simple explanation of the project for presenting to a viva examiner.
Usage: python scripts/build_viva_notes.py"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
D, S = ROOT / "docs" / "diagrams", ROOT / "docs" / "screenshots"
NAVY = RGBColor(0x1E, 0x3A, 0x5F)

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(1)
sec.top_margin = sec.bottom_margin = Inches(0.8)
doc.styles["Normal"].font.name, doc.styles["Normal"].font.size = "Calibri", Pt(12)
for n, sz in (("Heading 1", 18), ("Heading 2", 14)):
    h = doc.styles[n]
    h.font.name, h.font.size, h.font.color.rgb = "Calibri", Pt(sz), NAVY


def shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_)
    tcPr.append(shd)


def H1(t, brk=True):
    if brk:
        doc.add_page_break()
    doc.add_heading(t, 1)


def P(t, bold=False, size=None, align=None, after=6, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(t)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if align == "c":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)


def B(items, style="List Bullet"):
    for t in items:
        p = doc.add_paragraph(style=style)
        if t.startswith("**") and "**" in t[2:]:
            head, rest = t[2:].split("**", 1)
            p.add_run(head).bold = True
            p.add_run(rest)
        else:
            p.add_run(t)


def SAY(t):
    """A sentence you can say out loud to the examiner."""
    tb = doc.add_table(rows=1, cols=1)
    tb.style = "Table Grid"
    c = tb.rows[0].cells[0]
    shade(c, "EAF4EA")
    c.text = ""
    r = c.paragraphs[0].add_run("You can say:  "); r.bold = True
    c.paragraphs[0].add_run(t).italic = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def FIG(path, cap, w=6.2):
    doc.add_picture(str(path), width=Inches(w))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(cap, italic=True, size=10, align="c", after=8)


def TABLE(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "1E3A5F")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for c, w in zip(row.cells, widths):
            c.width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


# ----------------------------------------------------------------------------- cover
P("VIVA NOTES", True, 13, align="c", after=4)
P("CURIO", True, 36, align="c", after=0)
P("A luxury shopping website that suggests products and explains why", italic=True, size=14, align="c", after=24)
P("Short, simple notes to understand the project from zero and explain it to the examiner.", size=12, align="c", after=24)
TABLE(["Read in this order", "What you get"],
      [["1  The idea in 30 seconds", "What to say first"], ["2  How a website works (3 ideas)", "The basics, no jargon"], ["3  Our parts and how they connect", "The whole project on one picture"],
       ["4  Walk through one click", "The easiest way to explain everything"], ["5  The two suggestion features", "The main part of the project, in plain words"],
       ["6  Safety: passwords, stock, payments", "Questions examiners love"], ["7  How to run and show it", "Demo steps"], ["8  Quick answers", "20 likely questions with 1-2 line answers"]], [3.0, 3.5])

# ----------------------------------------------------------------------------- 1
H1("1  The idea in 30 seconds")
P("CURIO is an online shop for luxury items such as watches, handbags and perfumes. It works like a real shop: you can search, add things to a bag, pay and track the order. "
  "The special part is that when you open a product, the website suggests other products and tells you why.")
SAY("Our project is a luxury shopping website. Its main feature is a recommendation system. When you open a product, it shows similar products and products that other customers bought together with it, and it explains each suggestion.")
B(["**Built with:** Python, Streamlit (for the web pages), MongoDB (the database), scikit-learn (for the similarity maths).",
   "**What is real:** real brands and models, real photos, real Wikipedia descriptions for 30 products.",
   "**What is pretend:** the payment (a safe copy of Razorpay's test mode), delivery and tax. No money moves.",
   "**Size:** 125 products from 72 brands, 8 collections, 158 automated tests."])
FIG(S / "11_product_page.png", "The product page. The suggestions are shown below it.", 5.4)

# ----------------------------------------------------------------------------- 2
H1("2  How a website works (3 ideas)")
P("You only need three ideas to understand any website.")
B(["**Idea 1: Browser and server.** The browser (Chrome) is what you look at. The server is a program that decides what to show. When you click, the browser asks the server, and the server answers with a new page.",
   "**Idea 2: The server needs a place to keep data.** That place is the database. It stores users, products and orders so they are not lost when the program stops.",
   "**Idea 3: The rules live in the server.** Things like 'you cannot buy more than is in stock' are checked on the server, because a user could cheat in the browser."])
SAY("A website has a browser the user sees, a server that applies the rules, and a database that stores the data. In our project the server is a Python program and the database is MongoDB.")
P("Four words you will hear:", bold=True)
TABLE(["Word", "Meaning in plain words"],
      [["localhost:8501", "Our website running on our own computer, on door number 8501."], ["Streamlit", "A tool that lets us build the web pages using only Python."],
       ["MongoDB", "A database that stores each record as a flexible document, like a small form."], ["Docker", "A box that runs MongoDB for us, so we do not install it by hand."]], [1.8, 4.7])

# ----------------------------------------------------------------------------- 3
H1("3  Our parts and how they connect")
P("The project has three layers. Pages only talk to services. Services are the only ones who talk to the database.")
FIG(D / "14_tech_stack.png", "The layers, top to bottom.", 6.0)
TABLE(["Folder", "Job", "Example"],
      [["ui/", "What you see (the pages)", "Home page, product page, bag, admin"], ["shop/", "What the shop does (the rules)", "Add to bag, check coupon, pay, reduce stock"],
       ["MongoDB", "Where data is stored", "users, products, orders, payments ..."], ["tests/", "Checks that everything works", "158 automatic tests"]], [1.3, 2.5, 2.7])
SAY("We separated the code into pages and rules. The pages only show things; the rules, like stock and payment checks, are in one place, so every page is protected the same way and we can test the rules without a browser.")

# ----------------------------------------------------------------------------- 4
H1("4  Walk through one click")
P("The easiest way to explain the whole project is to follow one click. Say these steps in order.")
B(["You click a product picture.", "The page remembers 'show the product page for this product'. (Streamlit keeps this memory in something called session state.)",
   "The server asks the database for that product and shows its photo, price and description.",
   "It asks the first suggestion feature for 'You may also like'.", "It asks the second suggestion feature for 'Frequently bought together'.", "The page is drawn and sent to your browser."], "List Number")
FIG(D / "15_streamlit_rerun.png", "Streamlit runs the script again after every click, and the page redraws from the saved memory.", 6.0)
SAY("Streamlit runs the whole Python script again after each click. To remember what the user was doing, we store the current page and the logged-in user in session state.")
P("Now add to bag, checkout and pay:", bold=True)
B(["The bag stores product ids and quantities for that user.", "At checkout the server re-checks stock and reserves the pieces.", "Payment is verified on the server, then the order becomes 'Placed'.",
   "An admin moves it to Packed, Shipped and Delivered. The customer sees these steps on the tracking page."])

# ----------------------------------------------------------------------------- 5
H1("5  The two suggestion features")
P("This is the main part. Explain it slowly.")
doc.add_heading("Feature 1: You may also like (similar products)", 2)
P("Idea: two products are similar if their descriptions use similar words. A computer cannot read, so we turn each description into numbers, then compare the numbers.")
B(["Step 1: join the name, brand, category, tags and description into one text.", "Step 2 (TF-IDF): count the words. Words that appear in every product (like 'the') get a small weight. Rare words (like 'chronograph') get a big weight.",
   "Step 3 (cosine similarity): compare two products. The score is between 0 (nothing in common) and 1 (same words).",
   "Step 4: remove the product itself and sold-out items, and show the best matches with the reason (same brand, same collection, shared tags)."], "List Number")
FIG(D / "17_tfidf_example.png", "A tiny example: the two watches share the word 'watch', so they are similar; the perfume shares nothing.", 6.0)
SAY("We use TF-IDF to convert product text into numbers and cosine similarity to compare them. It is content-based, so it works even when we have no customer history.")
doc.add_heading("Feature 2: Frequently bought together", 2)
P("Idea: look at past orders and see which products often appear in the same order.")
B(["**Confidence:** of the orders with product A, what share also contain B. Example: 18 of 40 orders = 45%.",
   "**Lift:** is that more than chance? If B is in 30% of all orders, 45% / 30% = 1.5. Above 1 means they really go together.",
   "We show a pair only if it appears in at least 2 orders and has lift above 1."])
FIG(D / "18_basket_maths.png", "The same example in numbers.", 6.0)
SAY("A new shop has no orders, so we generated about 850 past orders from shopping patterns and marked them synthetic. Real orders placed on the site are added to the same pool.")
P("Be ready to say clearly: the starting orders are generated, not real customers.", bold=True)
P("How we checked quality: for each product, are the top 5 suggestions from the same subcategory? Similar-product suggestions score 0.59 against 0.40 for a simple 'same category' rule and 0.04 for random. This is a rough proxy, not proof of sales.")

# ----------------------------------------------------------------------------- 6
H1("6  Safety: passwords, stock and payments")
doc.add_heading("Passwords", 2)
P("We never store the password. We store a hash made with scrypt and a random salt. A hash is a one-way scramble: we can check a password against it, but we cannot turn it back into the password.")
doc.add_heading("Stock (what if two people buy the last item?)", 2)
P("When an order is placed, the database only subtracts stock if enough is left, in one atomic step. If any item in the order fails, the earlier items are put back. So stock never goes below zero.")
doc.add_heading("Payments", 2)
P("Our payment gateway is a safe copy of Razorpay's test mode. After you pay, the gateway gives a payment id and a signature. The server calculates the signature again with a secret key and compares. If someone changed the data in the browser, the signatures differ and the payment is rejected. The card number and CVV are never saved.")
FIG(D / "19_payment_signature.png", "Why the server checks the signature.", 6.0)
doc.add_heading("Other checks", 2)
B(["Every input is checked (email format, phone number, PIN code, card number checksum, quantity 1 to 10).", "Only admins can change order status; this is checked inside the rules, not only on the page.",
   "A user can only open their own orders.", "Unpaid online orders are cancelled after 15 minutes and the stock is returned."])

# ----------------------------------------------------------------------------- 7
H1("7  How to run and show it")
P("In a terminal, inside the project folder:")
tb = doc.add_table(rows=1, cols=1)
tb.style = "Table Grid"
c = tb.rows[0].cells[0]
shade(c, "F3F4F6")
c.text = ""
for i, line in enumerate(["./start.sh        # checks everything, starts MongoDB and the website", "open http://localhost:8501", "./stop.sh         # stops everything"]):
    para = c.paragraphs[0] if i == 0 else c.add_paragraph()
    r = para.add_run(line); r.font.name, r.font.size = "Consolas", Pt(10.5)
doc.add_paragraph()
P("Logins:", bold=True)
TABLE(["Role", "Email", "Password"], [["Customer", "demo@example.com", "Demo@1234"], ["Admin", "admin@curio.test", "Admin@1234"]], [1.5, 2.6, 2.4])
P("Demo script (3 minutes):", bold=True)
B(["Open the home page and scroll: the big word stays fixed and the pictures unfold as you scroll.", "Click Watches, open the Rolex Submariner. Point to the two suggestion sections and read one reason.",
   "Log in as the customer, press 'Add all to bag' under Frequently bought together, go to checkout.", "Pay with UPI success@razorpay. Show the order tracking page.",
   "Log in as admin and show the dashboard, then mark the order Packed.", "Mention: test payment values are on the Help page; 158 tests pass."], "List Number")
FIG(S / "24_admin_dashboard.png", "The admin dashboard.", 5.6)

# ----------------------------------------------------------------------------- 8
H1("8  Quick answers (20 likely questions)")
QA = [
 ("What is your project?", "A luxury e-commerce website that recommends similar products and products bought together, and explains why."),
 ("Which technologies?", "Python, Streamlit, MongoDB with PyMongo, scikit-learn, Pandas, Docker."),
 ("Why Streamlit?", "It builds web pages from Python, so one language does the pages, the rules and the maths."),
 ("Why MongoDB?", "Products have different attributes and an order contains its own lines, which suits documents. It was also asked for in the synopsis."),
 ("What is TF-IDF?", "A way to turn text into numbers where rare, meaningful words get more weight than common words."),
 ("What is cosine similarity?", "A score from 0 to 1 that says how similar two lists of numbers are."),
 ("Is it machine learning?", "It is a classical information-retrieval method and association analysis. Nothing is trained with labels and it is not deep learning."),
 ("What is confidence and lift?", "Confidence is how often B appears with A. Lift checks that this is more than chance; above 1 means a real link."),
 ("Where does the order data come from?", "About 850 orders are generated and marked synthetic. Real orders are added on top."),
 ("Is the payment real?", "No. It is a safe simulation of Razorpay test mode. No money moves and no card is stored."),
 ("How do you keep payments safe?", "The server checks a signature made with a secret key, and checks the amount, before marking the order paid."),
 ("How are passwords stored?", "As a scrypt hash with a random salt, never as plain text."),
 ("How do you stop selling the last item twice?", "The stock update only happens if enough stock remains, and it is atomic. Failures are rolled back."),
 ("What happens if the user does not pay?", "After 15 minutes the order is cancelled and the stock is returned."),
 ("How does the page remember the user?", "Streamlit session state, a small memory kept for each browser tab."),
 ("Are the photos and descriptions real?", "Photos are real Unsplash photos; 30 descriptions are Wikipedia text and the rest are written by us. Some photos show a similar piece, and prices are indicative."),
 ("How did you test it?", "158 automatic tests for the rules and the pages, also run on a real MongoDB, plus a run in a real browser with real clicks."),
 ("What are the limitations?", "Simulated payment and delivery, generated starting orders, suggestions are not personal to the user."),
 ("What could be added later?", "Real Razorpay, personal recommendations from each user's history, image-based similarity, cloud hosting."),
 ("Which part are you most proud of?", "The two explainable suggestion features and the server-side payment check."),
]
for i, (q, a) in enumerate(QA, 1):
    p = doc.add_paragraph()
    p.add_run(f"{i}. {q}").bold = True
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    doc.add_paragraph(a).paragraph_format.space_after = Pt(5)
P("For deeper answers, see Complete_Guide.docx (chapter 19 has 30 detailed questions).", italic=True)

out = ROOT / "docs" / "Viva_Notes.docx"
doc.save(out)
print("saved", out)
